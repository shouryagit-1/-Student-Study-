# -*- coding: utf-8 -*-
import os
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
from reportlab.pdfgen import canvas

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_border()
            super().showPage()
        super().save()

    def draw_page_border(self):
        self.saveState()
        w, h = self._pagesize
        margin = 32
        self.setStrokeColor(colors.HexColor("#222222"))
        self.setLineWidth(0.8)
        self.rect(margin, margin, w - 2 * margin, h - 2 * margin, stroke=1, fill=0)
        self.restoreState()

def build_pdf(filename="AI_Phishing_Detection_Synopsis_Shresthi_Kesharwani.pdf"):
    doc = SimpleDocTemplate(
        filename,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=50,
        bottomMargin=50
    )
    
    styles = getSampleStyleSheet()
    
    t_center_bold_lg = ParagraphStyle(
        'TCenterBoldLg',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=17,
        leading=22,
        alignment=1,
        textColor=colors.black,
        spaceAfter=14
    )
    
    t_center_bold_md = ParagraphStyle(
        'TCenterBoldMd',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=15,
        leading=20,
        alignment=1,
        textColor=colors.black,
        spaceAfter=10
    )
    
    t_center_regular = ParagraphStyle(
        'TCenterReg',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=12.5,
        leading=18,
        alignment=1,
        textColor=colors.black,
        spaceAfter=4
    )
    
    t_center_italic = ParagraphStyle(
        'TCenterItalic',
        parent=styles['Normal'],
        fontName='Times-Italic',
        fontSize=12,
        leading=17,
        alignment=1,
        textColor=colors.black,
        spaceAfter=4
    )
    
    h1_style = ParagraphStyle(
        'H1Style',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=14,
        leading=19,
        textColor=colors.black,
        spaceBefore=8,
        spaceAfter=14
    )
    
    h2_style = ParagraphStyle(
        'H2Style',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=12.5,
        leading=17,
        textColor=colors.black,
        spaceBefore=12,
        spaceAfter=6
    )
    
    body_style = ParagraphStyle(
        'BodyJustified',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=11.5,
        leading=17.5,
        alignment=4,
        textColor=colors.black,
        spaceAfter=12
    )
    
    bullet_style = ParagraphStyle(
        'BulletStyle',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=11.5,
        leading=17,
        alignment=4,
        textColor=colors.black,
        leftIndent=24,
        firstLineIndent=-14,
        spaceAfter=8
    )
    
    caption_style = ParagraphStyle(
        'CaptionStyle',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=10,
        leading=14,
        alignment=1,
        textColor=colors.black,
        spaceBefore=6,
        spaceAfter=6
    )
    
    subcaption_style = ParagraphStyle(
        'SubCaptionStyle',
        parent=styles['Normal'],
        fontName='Times-Italic',
        fontSize=8.5,
        leading=12,
        alignment=1,
        textColor=colors.HexColor("#333333"),
        spaceBefore=4,
        spaceAfter=14
    )
    
    ref_style = ParagraphStyle(
        'RefStyle',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=11,
        leading=16.5,
        alignment=4,
        leftIndent=28,
        firstLineIndent=-28,
        textColor=colors.black,
        spaceAfter=14
    )

    story = []
    
    # ==================== PAGE 1: TITLE PAGE ====================
    story.append(Spacer(1, 24))
    story.append(Paragraph("<b>SYNOPSIS</b>", t_center_bold_lg))
    story.append(Paragraph("<b>ON</b>", t_center_bold_md))
    story.append(Paragraph("<b>&ldquo;AI BASED PHISHING AND SCAM DETECTION SYSTEM&rdquo;</b>", t_center_bold_lg))
    
    story.append(Spacer(1, 20))
    story.append(Paragraph("Submitted in", t_center_regular))
    story.append(Paragraph("Partial Fulfillment of the requirements for the Award of the Degree", t_center_italic))
    story.append(Paragraph("<i>of</i>", t_center_italic))
    story.append(Paragraph("<b>Bachelor of Technology</b>", t_center_regular))
    story.append(Paragraph("<i>In</i>", t_center_italic))
    story.append(Paragraph("<b>Computer Science and Engineering</b>", t_center_regular))
    story.append(Paragraph("<i>by</i>", t_center_italic))
    
    story.append(Spacer(1, 14))
    story.append(Paragraph("<b>Student Name: Shresthi Kesharwani</b>", t_center_regular))
    story.append(Paragraph("<b>Roll No.: 2501641520196</b>", t_center_regular))
    
    story.append(Spacer(1, 16))
    story.append(Paragraph("Under the supervision of", t_center_regular))
    story.append(Paragraph("<b>Supervisor Name: Mr. Rajesh Kumar</b>", t_center_regular))
    story.append(Paragraph("(Assistant Professor)", t_center_regular))
    
    story.append(Spacer(1, 18))
    if os.path.exists("psit_logo.png"):
        story.append(Image("psit_logo.png", width=1.9*72, height=0.75*72))
    story.append(Spacer(1, 14))
    story.append(Paragraph("<b>Pranveer Singh Institute of Technology.</b>", t_center_bold_md))
    story.append(Paragraph("Kanpur - Agra - Delhi National Highway - 19", t_center_regular))
    story.append(Paragraph("Bhauti - Kanpur - 209305.", t_center_regular))
    
    story.append(PageBreak())
    
    # ==================== PAGE 2: INTRODUCTION ====================
    story.append(Paragraph("<b>1. Introduction</b>", h1_style))
    story.append(Paragraph(
        "Phishing and online scams are among the most common and damaging cyber threats today. "
        "In a phishing attack, an attacker imitates a trusted bank, company or service through a fake website, "
        "email or SMS, and tricks the victim into revealing passwords, card numbers, OTPs or other sensitive data. "
        "Scams such as fake job offers, lottery messages, fake payment links and impersonation calls follow a "
        "similar pattern and cause large financial losses every year, particularly in India where digital payments "
        "are growing quickly.",
        body_style
    ))
    story.append(Paragraph(
        "Traditional defences such as blacklists and fixed rule-based filters only block threats that are already known. "
        "Attackers register new domains every day and change their messages continuously, so these methods often "
        "fail against newly created (&ldquo;zero-day&rdquo;) phishing pages. Artificial Intelligence (AI) and Machine Learning (ML) "
        "can address this gap by learning the patterns that separate malicious content from genuine content, so that "
        "unseen threats can also be recognised.",
        body_style
    ))
    story.append(Paragraph(
        "This project proposes an AI Based Phishing and Scam Detection System that analyses a given URL, email or "
        "message text and classifies it as Safe, Suspicious or Phishing/Scam, together with a risk score and a "
        "short explanation. The system extracts URL-based features (length, number of dots and special characters, "
        "use of IP address, HTTPS status, suspicious keywords), domain-based features (domain age, WHOIS details) "
        "and content-based features (urgency words, requests for credentials, mismatched links). A trained classifier "
        "then predicts the result in real time through a web application and an optional browser extension.",
        body_style
    ))
    story.append(Paragraph(
        "<b>Special terms:</b> Phishing &ndash; fraudulent attempt to steal sensitive information by impersonating a "
        "trusted entity; URL &ndash; Uniform Resource Locator, the web address of a resource; Feature Extraction &ndash; "
        "converting raw data into measurable attributes; Classifier &ndash; an ML model that assigns an input to a "
        "class; NLP &ndash; Natural Language Processing, used for analysing message text; Zero-day attack &ndash; an attack "
        "not yet present in any blacklist.",
        body_style
    ))
    
    story.append(PageBreak())
    
    # ==================== PAGE 3: PROJECT OBJECTIVE ====================
    story.append(Paragraph("<b>2. Project Objective</b>", h1_style))
    story.append(Paragraph(
        "The main objective of this project is to design and develop an intelligent, easy-to-use system that "
        "detects phishing websites and scam messages accurately and in real time, and warns the user before "
        "any harm is done. The specific objectives are:",
        body_style
    ))
    
    objectives = [
        "To collect and pre-process a labelled dataset of phishing and legitimate URLs and scam and genuine messages from sources such as PhishTank, the UCI Phishing Websites dataset and Kaggle.",
        "To extract meaningful URL, domain and text features that help distinguish malicious content from genuine content.",
        "To train and compare ML models (Logistic Regression, Random Forest, XGBoost and an NLP-based text classifier) and select the one with the best accuracy, precision, recall and F1-score.",
        "To build a web application and browser extension where a user can submit a link or message and instantly get the result with a risk score and the reasons behind it.",
        "To reduce false alarms and to detect newly created phishing links that are not yet in any blacklist.",
        "To store scan history and generate simple reports for analysis and awareness."
    ]
    
    for obj in objectives:
        story.append(Paragraph(f"&bull; &nbsp; {obj}", bullet_style))
        
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "At the end of the project, ordinary users, students and small businesses will have a simple and free tool "
        "that helps them identify dangerous links and messages before clicking or replying, which reduces the risk "
        "of identity theft and financial fraud.",
        body_style
    ))
    
    story.append(PageBreak())
    
    # ==================== PAGE 4: FEASIBILITY STUDY ====================
    story.append(Paragraph("<b>3. Feasibility Study</b>", h1_style))
    story.append(Paragraph(
        "A feasibility study checks whether the proposed system can be built successfully with the available "
        "time, skills and resources. The project was examined from the following points of view.",
        body_style
    ))
    
    story.append(Paragraph(
        "<b>Technical Feasibility:</b> The project uses open-source, well-documented technologies such as Python, "
        "scikit-learn, Flask and JavaScript. Public datasets are available for training, and the model can run on "
        "an ordinary laptop. The team has the required knowledge of programming and basic machine learning, so the "
        "project is technically feasible.",
        body_style
    ))
    
    story.append(Paragraph(
        "<b>Operational Feasibility:</b> The system has a simple interface where the user only pastes a URL or message, "
        "or uses the browser extension, and gets a clear result. It needs no special training and fits easily into the "
        "way people already browse and use email, so it is operationally feasible.",
        body_style
    ))
    
    story.append(Paragraph(
        "<b>Economic Feasibility:</b> All major tools and libraries are free and open source and the datasets are publicly "
        "available. No costly hardware is needed, so the development cost is very low and the project is economically "
        "feasible for a student mini project.",
        body_style
    ))
    
    story.append(Paragraph(
        "<b>Schedule Feasibility:</b> The work is divided into requirement analysis, data collection, model training, "
        "application development, testing and documentation. These phases can be completed within the semester time frame.",
        body_style
    ))
    
    story.append(Paragraph(
        "<b>Legal Feasibility:</b> Only publicly available datasets are used and no private user data is collected "
        "without consent. The system is designed for defence and awareness only and follows ethical and IT-law guidelines, "
        "so there is no legal obstacle.",
        body_style
    ))
    
    story.append(PageBreak())
    
    # ==================== PAGE 5: METHODOLOGY / PLANNING OF WORK ====================
    story.append(Paragraph("<b>4. Methodology / Planning of Work</b>", h1_style))
    story.append(Paragraph(
        "The project will follow a step-by-step approach, from data collection to deployment:",
        body_style
    ))
    
    steps = [
        "<b>Requirement analysis:</b> study existing phishing and scam techniques, detection methods and their limitations.",
        "<b>Data collection:</b> gather labelled phishing and legitimate URLs and scam and genuine messages from public sources.",
        "<b>Data pre-processing:</b> remove duplicates and missing values, normalise data and balance the classes.",
        "<b>Feature extraction:</b> compute URL, domain and content features; use TF-IDF for message text.",
        "<b>Model training and evaluation:</b> train several models, tune hyper-parameters and compare them using accuracy, precision, recall, F1-score and the confusion matrix.",
        "<b>Application development:</b> integrate the best model into a Flask back end with a web interface and a browser extension.",
        "<b>Testing:</b> perform unit, integration and real-world testing with fresh phishing links and check false-positive rates.",
        "<b>Deployment and documentation:</b> host the application locally or on a free cloud platform and prepare the final report."
    ]
    
    for s in steps:
        story.append(Paragraph(f"&bull; &nbsp; {s}", bullet_style))
        
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        "The flow of information between the modules is shown in the architecture diagram below.",
        body_style
    ))
    
    story.append(PageBreak())
    
    # ==================== PAGE 6: ARCHITECTURE & TOOLS ====================
    story.append(Paragraph("<b>Fig. 1: System Architecture of the Proposed Phishing &amp; Scam Detection System</b>", caption_style))
    
    if os.path.exists("architecture_diagram.png"):
        story.append(Image("architecture_diagram.png", width=6.6*72, height=2.6*72))
    
    story.append(Paragraph("<i>Training pipeline: Labelled dataset (PhishTank / UCI / Kaggle) &rarr; Cleaning &rarr; Feature selection &rarr; Model training &amp; evaluation &rarr; Saved model</i>", subcaption_style))
    
    story.append(Paragraph("<b>5. Tools / Technology Used</b>", h1_style))
    story.append(Paragraph("<b>5.1 Minimum Hardware Requirements</b>", h2_style))
    story.append(Paragraph("Hardware required for the development of the project.", body_style))
    
    hw_items = [
        "CPU: Intel Core i5 (8th Gen) or AMD Ryzen 5 or higher, 2.4 GHz or above",
        "RAM: 8 GB (16 GB recommended for model training)",
        "GPU: Not mandatory; integrated graphics are sufficient. An NVIDIA GTX 1650 or better is optional for deep-learning experiments",
        "HDD: 256 GB SSD or 500 GB HDD with at least 20 GB free space",
        "Others (if any): Stable internet connection (for dataset download and domain look-ups), standard keyboard, mouse and monitor"
    ]
    for item in hw_items:
        story.append(Paragraph(f"&bull; &nbsp; {item}", bullet_style))
        
    story.append(Paragraph("<b>5.2 Minimum Software Requirements</b>", h2_style))
    story.append(Paragraph("Software required for the development of the project.", body_style))
    
    sw_items = [
        "OS: Windows 10/11 (64-bit) or Ubuntu 22.04 LTS",
        "Programming languages: Python 3.10 or above; HTML, CSS and JavaScript for the front end and browser extension",
        "Libraries and frameworks: Pandas, NumPy, scikit-learn, XGBoost, NLTK, tldextract, python-whois, Matplotlib, Flask",
        "Database: SQLite for storing scan history and logs",
        "Other tools: Jupyter Notebook / Google Colab, Visual Studio Code, Git and GitHub, Google Chrome (for extension testing)"
    ]
    for item in sw_items:
        story.append(Paragraph(f"&bull; &nbsp; {item}", bullet_style))
        
    story.append(PageBreak())
    
    # ==================== PAGE 7: REFERENCES ====================
    story.append(Paragraph("<b>6. References</b>", h1_style))
    story.append(Spacer(1, 6))
    
    refs = [
        "[1] A. Basit, M. Zafar, X. Liu, A. R. Javed, Z. Jalil, and K. Kifayat, \u201cA comprehensive survey of AI-enabled phishing attacks detection techniques,\u201d <i>Telecommunication Systems</i>, vol. 76, no. 1, pp. 139\u2013154, 2021.",
        "[2] R. M. Mohammad, F. Thabtah, and L. McCluskey, \u201cPredicting phishing websites based on self-structuring neural network,\u201d <i>Neural Computing and Applications</i>, vol. 25, no. 2, pp. 443\u2013458, 2014.",
        "[3] F. Pedregosa et al., \u201cScikit-learn: Machine learning in Python,\u201d <i>Journal of Machine Learning Research</i>, vol. 12, pp. 2825\u20132830, 2011.",
        "[4] T. Chen and C. Guestrin, \u201cXGBoost: A scalable tree boosting system,\u201d in <i>Proc. 22nd ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining</i>, San Francisco, CA, USA, 2016, pp. 785\u2013794.",
        "[5] PhishTank, \u201cPhishTank \u2013 Join the fight against phishing.\u201d [Online]. Available: https://www.phishtank.org",
        "[6] Anti-Phishing Working Group (APWG), \u201cPhishing Activity Trends Reports.\u201d [Online]. Available: https://apwg.org"
    ]
    
    for ref in refs:
        story.append(Paragraph(ref, ref_style))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"Successfully generated PDF: {filename}")


def set_docx_page_border(section):
    sectPr = section._sectPr
    pgBorders = parse_xml(
        r'<w:pgBorders %s w:offsetFrom="page">'
        r'  <w:top w:val="single" w:sz="6" w:space="24" w:color="333333"/>'
        r'  <w:left w:val="single" w:sz="6" w:space="24" w:color="333333"/>'
        r'  <w:bottom w:val="single" w:sz="6" w:space="24" w:color="333333"/>'
        r'  <w:right w:val="single" w:sz="6" w:space="24" w:color="333333"/>'
        r'</w:pgBorders>' % nsdecls('w')
    )
    sectPr.append(pgBorders)

def build_docx(filename="AI_Phishing_Detection_Synopsis_Shresthi_Kesharwani.docx"):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        set_docx_page_border(section)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    # --- PAGE 1: COVER ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("SYNOPSIS\nON\n")
    run.bold = True
    run.font.size = Pt(16)
    run2 = p.add_run("\u201cAI BASED PHISHING AND SCAM DETECTION SYSTEM\u201d")
    run2.bold = True
    run2.font.size = Pt(16)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(16)
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.line_spacing = 1.3
    r = p2.add_run("Submitted in\n")
    r.font.size = Pt(12)
    r = p2.add_run("Partial Fulfillment of the requirements for the Award of the Degree\nof\n")
    r.italic = True
    r.font.size = Pt(12)
    r = p2.add_run("Bachelor of Technology\n")
    r.bold = True
    r.font.size = Pt(13)
    r = p2.add_run("In\n")
    r.italic = True
    r.font.size = Pt(12)
    r = p2.add_run("Computer Science and Engineering\n")
    r.bold = True
    r.font.size = Pt(13)
    r = p2.add_run("by\n")
    r.italic = True
    r.font.size = Pt(12)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(14)
    p3.paragraph_format.space_after = Pt(16)
    r = p3.add_run("Student Name: Shresthi Kesharwani\nRoll No.: 2501641520196")
    r.bold = True
    r.font.size = Pt(13)

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(12)
    p4.paragraph_format.space_after = Pt(16)
    r = p4.add_run("Under the supervision of\n")
    r.font.size = Pt(12)
    r = p4.add_run("Supervisor Name: Mr. Rajesh Kumar\n")
    r.bold = True
    r.font.size = Pt(13)
    r = p4.add_run("(Assistant Professor)")
    r.font.size = Pt(12)

    if os.path.exists("psit_logo.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(10)
        p_img.add_run().add_picture("psit_logo.png", width=Inches(2.0))

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(8)
    p5.paragraph_format.space_after = Pt(4)
    r = p5.add_run("Pranveer Singh Institute of Technology.\n")
    r.bold = True
    r.font.size = Pt(14)
    r = p5.add_run("Kanpur - Agra - Delhi National Highway - 19\nBhauti - Kanpur - 209305.")
    r.font.size = Pt(12)

    doc.add_page_break()

    # --- PAGE 2: INTRODUCTION ---
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("1. Introduction")
    r.bold = True
    r.font.size = Pt(14)

    intro_paras = [
        "Phishing and online scams are among the most common and damaging cyber threats today. In a phishing attack, an attacker imitates a trusted bank, company or service through a fake website, email or SMS, and tricks the victim into revealing passwords, card numbers, OTPs or other sensitive data. Scams such as fake job offers, lottery messages, fake payment links and impersonation calls follow a similar pattern and cause large financial losses every year, particularly in India where digital payments are growing quickly.",
        "Traditional defences such as blacklists and fixed rule-based filters only block threats that are already known. Attackers register new domains every day and change their messages continuously, so these methods often fail against newly created (\"zero-day\") phishing pages. Artificial Intelligence (AI) and Machine Learning (ML) can address this gap by learning the patterns that separate malicious content from genuine content, so that unseen threats can also be recognised.",
        "This project proposes an AI Based Phishing and Scam Detection System that analyses a given URL, email or message text and classifies it as Safe, Suspicious or Phishing/Scam, together with a risk score and a short explanation. The system extracts URL-based features (length, number of dots and special characters, use of IP address, HTTPS status, suspicious keywords), domain-based features (domain age, WHOIS details) and content-based features (urgency words, requests for credentials, mismatched links). A trained classifier then predicts the result in real time through a web application and an optional browser extension.",
        "Special terms: Phishing \u2013 fraudulent attempt to steal sensitive information by impersonating a trusted entity; URL \u2013 Uniform Resource Locator, the web address of a resource; Feature Extraction \u2013 converting raw data into measurable attributes; Classifier \u2013 an ML model that assigns an input to a class; NLP \u2013 Natural Language Processing, used for analysing message text; Zero-day attack \u2013 an attack not yet present in any blacklist."
    ]

    for text in intro_paras:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.25
        if text.startswith("Special terms:"):
            r_bold = p.add_run("Special terms: ")
            r_bold.bold = True
            p.add_run(text[len("Special terms: "):])
        else:
            p.add_run(text)

    doc.add_page_break()

    # --- PAGE 3: PROJECT OBJECTIVE ---
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("2. Project Objective")
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.25
    p.add_run("The main objective of this project is to design and develop an intelligent, easy-to-use system that detects phishing websites and scam messages accurately and in real time, and warns the user before any harm is done. The specific objectives are:")

    objectives = [
        "To collect and pre-process a labelled dataset of phishing and legitimate URLs and scam and genuine messages from sources such as PhishTank, the UCI Phishing Websites dataset and Kaggle.",
        "To extract meaningful URL, domain and text features that help distinguish malicious content from genuine content.",
        "To train and compare ML models (Logistic Regression, Random Forest, XGBoost and an NLP-based text classifier) and select the one with the best accuracy, precision, recall and F1-score.",
        "To build a web application and browser extension where a user can submit a link or message and instantly get the result with a risk score and the reasons behind it.",
        "To reduce false alarms and to detect newly created phishing links that are not yet in any blacklist.",
        "To store scan history and generate simple reports for analysis and awareness."
    ]

    for obj in objectives:
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.2
        p.add_run(obj)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.25
    p.add_run("At the end of the project, ordinary users, students and small businesses will have a simple and free tool that helps them identify dangerous links and messages before clicking or replying, which reduces the risk of identity theft and financial fraud.")

    doc.add_page_break()

    # --- PAGE 4: FEASIBILITY STUDY ---
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("3. Feasibility Study")
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.25
    p.add_run("A feasibility study checks whether the proposed system can be built successfully with the available time, skills and resources. The project was examined from the following points of view.")

    feas_points = [
        ("Technical Feasibility: ", "The project uses open-source, well-documented technologies such as Python, scikit-learn, Flask and JavaScript. Public datasets are available for training, and the model can run on an ordinary laptop. The team has the required knowledge of programming and basic machine learning, so the project is technically feasible."),
        ("Operational Feasibility: ", "The system has a simple interface where the user only pastes a URL or message, or uses the browser extension, and gets a clear result. It needs no special training and fits easily into the way people already browse and use email, so it is operationally feasible."),
        ("Economic Feasibility: ", "All major tools and libraries are free and open source and the datasets are publicly available. No costly hardware is needed, so the development cost is very low and the project is economically feasible for a student mini project."),
        ("Schedule Feasibility: ", "The work is divided into requirement analysis, data collection, model training, application development, testing and documentation. These phases can be completed within the semester time frame."),
        ("Legal Feasibility: ", "Only publicly available datasets are used and no private user data is collected without consent. The system is designed for defence and awareness only and follows ethical and IT-law guidelines, so there is no legal obstacle.")
    ]

    for title, desc in feas_points:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.25
        r_title = p.add_run(title)
        r_title.bold = True
        p.add_run(desc)

    doc.add_page_break()

    # --- PAGE 5: METHODOLOGY / PLANNING OF WORK ---
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("4. Methodology / Planning of Work")
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.25
    p.add_run("The project will follow a step-by-step approach, from data collection to deployment:")

    method_steps = [
        ("Requirement analysis: ", "study existing phishing and scam techniques, detection methods and their limitations."),
        ("Data collection: ", "gather labelled phishing and legitimate URLs and scam and genuine messages from public sources."),
        ("Data pre-processing: ", "remove duplicates and missing values, normalise data and balance the classes."),
        ("Feature extraction: ", "compute URL, domain and content features; use TF-IDF for message text."),
        ("Model training and evaluation: ", "train several models, tune hyper-parameters and compare them using accuracy, precision, recall, F1-score and the confusion matrix."),
        ("Application development: ", "integrate the best model into a Flask back end with a web interface and a browser extension."),
        ("Testing: ", "perform unit, integration and real-world testing with fresh phishing links and check false-positive rates."),
        ("Deployment and documentation: ", "host the application locally or on a free cloud platform and prepare the final report.")
    ]

    for title, desc in method_steps:
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(title)
        r.bold = True
        p.add_run(desc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.25
    p.add_run("The flow of information between the modules is shown in the architecture diagram below.")

    doc.add_page_break()

    # --- PAGE 6: ARCHITECTURE DIAGRAM & TOOLS ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Fig. 1: System Architecture of the Proposed Phishing & Scam Detection System")
    r.bold = True
    r.font.size = Pt(11)

    if os.path.exists("architecture_diagram.png"):
        p_diag = doc.add_paragraph()
        p_diag.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_diag.paragraph_format.space_after = Pt(4)
        p_diag.add_run().add_picture("architecture_diagram.png", width=Inches(6.2))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r = p_sub.add_run("Training pipeline: Labelled dataset (PhishTank / UCI / Kaggle) \u2192 Cleaning \u2192 Feature selection \u2192 Model training & evaluation \u2192 Saved model")
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(80, 80, 80)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("5. Tools / Technology Used")
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("5.1 Minimum Hardware Requirements")
    r.bold = True
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("Hardware required for the development of the project.")

    hw_items = [
        "CPU: Intel Core i5 (8th Gen) or AMD Ryzen 5 or higher, 2.4 GHz or above",
        "RAM: 8 GB (16 GB recommended for model training)",
        "GPU: Not mandatory; integrated graphics are sufficient. An NVIDIA GTX 1650 or better is optional for deep-learning experiments",
        "HDD: 256 GB SSD or 500 GB HDD with at least 20 GB free space",
        "Others (if any): Stable internet connection (for dataset download and domain look-ups), standard keyboard, mouse and monitor"
    ]
    for item in hw_items:
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.add_run(item)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("5.2 Minimum Software Requirements")
    r.bold = True
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("Software required for the development of the project.")

    sw_items = [
        "OS: Windows 10/11 (64-bit) or Ubuntu 22.04 LTS",
        "Programming languages: Python 3.10 or above; HTML, CSS and JavaScript for the front end and browser extension",
        "Libraries and frameworks: Pandas, NumPy, scikit-learn, XGBoost, NLTK, tldextract, python-whois, Matplotlib, Flask",
        "Database: SQLite for storing scan history and logs",
        "Other tools: Jupyter Notebook / Google Colab, Visual Studio Code, Git and GitHub, Google Chrome (for extension testing)"
    ]
    for item in sw_items:
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.add_run(item)

    doc.add_page_break()

    # --- PAGE 7: REFERENCES ---
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("6. References")
    r.bold = True
    r.font.size = Pt(14)

    refs_data = [
        ("[1] ", "A. Basit, M. Zafar, X. Liu, A. R. Javed, Z. Jalil, and K. Kifayat, \u201cA comprehensive survey of AI-enabled phishing attacks detection techniques,\u201d ", "Telecommunication Systems", ", vol. 76, no. 1, pp. 139\u2013154, 2021."),
        ("[2] ", "R. M. Mohammad, F. Thabtah, and L. McCluskey, \u201cPredicting phishing websites based on self-structuring neural network,\u201d ", "Neural Computing and Applications", ", vol. 25, no. 2, pp. 443\u2013458, 2014."),
        ("[3] ", "F. Pedregosa et al., \u201cScikit-learn: Machine learning in Python,\u201d ", "Journal of Machine Learning Research", ", vol. 12, pp. 2825\u20132830, 2011."),
        ("[4] ", "T. Chen and C. Guestrin, \u201cXGBoost: A scalable tree boosting system,\u201d in ", "Proc. 22nd ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining", ", San Francisco, CA, USA, 2016, pp. 785\u2013794."),
        ("[5] ", "PhishTank, \u201cPhishTank \u2013 Join the fight against phishing.\u201d [Online]. Available: https://www.phishtank.org", "", ""),
        ("[6] ", "Anti-Phishing Working Group (APWG), \u201cPhishing Activity Trends Reports.\u201d [Online]. Available: https://apwg.org", "", "")
    ]

    for num, text_pre, italic_text, text_post in refs_data:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        p.add_run(num)
        p.add_run(text_pre)
        if italic_text:
            r_it = p.add_run(italic_text)
            r_it.italic = True
        if text_post:
            p.add_run(text_post)

    doc.save(filename)
    print(f"Successfully generated DOCX: {filename}")

if __name__ == "__main__":
    build_pdf()
    build_docx()
